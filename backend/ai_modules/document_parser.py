"""
Document Parser — extracts text from PDF, DOCX, PPTX, TXT files.
"""
import os
import re


def parse_document(filepath: str) -> dict:
    """
    Parse a document and return extracted text with metadata.
    Supports: PDF, DOCX, PPTX, TXT
    """
    ext = os.path.splitext(filepath)[1].lower()
    text = ''
    metadata = {'pages': 0, 'word_count': 0}

    try:
        if ext == '.pdf':
            text, metadata = _parse_pdf(filepath)
        elif ext == '.docx':
            text, metadata = _parse_docx(filepath)
        elif ext in ('.pptx', '.ppt'):
            text, metadata = _parse_pptx(filepath)
        elif ext == '.txt':
            text, metadata = _parse_txt(filepath)
        else:
            raise ValueError(f"Unsupported file type: {ext}")
    except Exception as e:
        return {'success': False, 'error': str(e), 'text': '', 'metadata': metadata}

    text = clean_text(text)
    metadata['word_count'] = len(text.split())
    metadata['char_count'] = len(text)

    return {'success': True, 'text': text, 'metadata': metadata, 'extension': ext}


def _parse_pdf(filepath: str):
    import PyPDF2
    text_parts = []
    with open(filepath, 'rb') as f:
        reader = PyPDF2.PdfReader(f)
        pages = len(reader.pages)
        for page in reader.pages:
            text_parts.append(page.extract_text() or '')
    return '\n'.join(text_parts), {'pages': pages}


def _parse_docx(filepath: str):
    from docx import Document
    doc = Document(filepath)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Also extract from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)
    text = '\n'.join(paragraphs)
    return text, {'pages': len(doc.paragraphs) // 30 + 1}


def _parse_pptx(filepath: str):
    from pptx import Presentation
    prs = Presentation(filepath)
    slides_text = []
    for i, slide in enumerate(prs.slides):
        slide_texts = []
        for shape in slide.shapes:
            if hasattr(shape, 'text') and shape.text.strip():
                slide_texts.append(shape.text)
        slides_text.append(f"[Slide {i+1}]\n" + '\n'.join(slide_texts))
    text = '\n\n'.join(slides_text)
    return text, {'pages': len(prs.slides)}


def _parse_txt(filepath: str):
    with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
        text = f.read()
    lines = text.count('\n')
    return text, {'pages': lines // 40 + 1}


def clean_text(text: str) -> str:
    """Remove noise while preserving meaningful content."""
    text = re.sub(r'\s+', ' ', text)
    text = re.sub(r'[^\x00-\x7F]+', ' ', text)  # remove non-ASCII
    text = re.sub(r'\.{3,}', '.', text)
    text = re.sub(r'-{3,}', ' ', text)
    text = text.strip()
    return text


def chunk_text(text: str, chunk_size: int = 500, overlap: int = 100) -> list:
    """
    Split text into overlapping chunks for processing.
    Returns list of text chunks.
    """
    words = text.split()
    chunks = []
    step = chunk_size - overlap
    for i in range(0, len(words), step):
        chunk = ' '.join(words[i:i + chunk_size])
        if len(chunk.split()) >= 50:  # skip tiny chunks
            chunks.append(chunk)
    return chunks
